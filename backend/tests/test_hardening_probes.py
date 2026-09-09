"""Adversarial verification for Part 1; no hosted-provider claims are made here."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import io

import pytest
from pypdf import PdfReader
from docx import Document

from app.compliance.export import ExportContext, export_pdf, export_docx
from tests.test_export import _report, _flat


@pytest.mark.parametrize('index', range(9))
def test_each_registry_deletion_is_rejected(index, monkeypatch):
    from tests import test_regressions as registry
    key = sorted(registry.DEFECTS)[index]
    monkeypatch.delitem(registry.DEFECTS, key)
    with pytest.raises(AssertionError, match='unregistered defect'):
        registry.test_every_pinned_test_names_a_registered_defect()


def test_concurrent_unseen_reports_retain_all_requirements_in_both_formats():
    def render(index):
        report = _report().model_copy(deep=True)
        report.vendor_name = f'PROBE-VENDOR-{index}'
        original = report.items[0]
        report.items = [original.model_copy(update={'requirement': f'REQUIREMENT-{index}-{i:03d} supporting evidence'}, deep=True) for i in range(73)]
        context = ExportContext(generated_at=datetime(2026, 9, 5, tzinfo=timezone.utc))
        pdf = _flat(' '.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(export_pdf(report, context))).pages))
        doc = Document(io.BytesIO(export_docx(report, context)))
        docx = _flat(' '.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]))
        for text in (pdf, docx):
            assert report.vendor_name in text
            assert 'not a score' in text
            for i in range(73):
                assert f'REQUIREMENT-{index}-{i:03d}' in text
            for other in range(4):
                if other != index:
                    assert f'PROBE-VENDOR-{other}' not in text
        return True
    with ThreadPoolExecutor(max_workers=4) as pool:
        assert all(pool.map(render, range(4)))


def test_pdf_preserves_literal_markup_in_source_values():
    report = _report()
    report.vendor_name = 'Contractor <b>Literal</b> & Partners'
    text = _flat(' '.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(export_pdf(report))).pages))
    assert report.vendor_name in text


def test_token_budget_does_not_spend_the_same_refill_twice(monkeypatch):
    from app.llm import ratelimit
    clock = [0.0]
    monkeypatch.setattr(ratelimit.time, 'monotonic', lambda: clock[0])
    monkeypatch.setattr(ratelimit.time, 'sleep', lambda seconds: clock.__setitem__(0, clock[0] + seconds))
    limiter = ratelimit.RateLimiter(0, tokens_per_minute=600)
    limiter.acquire(600)
    limiter.acquire(20)
    assert clock[0] == 2
    limiter.acquire(20)
    assert clock[0] == 4, 'the refill spent by the previous waiter was spent again'
