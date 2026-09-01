"""Generates real PDF/DOCX files for testing the ingestion pipeline.

Parsers must be exercised against actual binary formats -- page breaks, tables,
headers -- not against clean strings. Section 9.2.2 step 5 makes the same point
about the synthetic vendor set, and this is the same machinery in miniature.
"""

from __future__ import annotations

from pathlib import Path

# Known ground truth for the generated notification, so parser and extractor
# tests can both assert against fixed values.
NOTIFICATION_TRUTH = {
    "tender_id": "TSTS/2026/IT/0042",
    "title": "Supply, Installation and Commissioning of Network Equipment",
    "issuing_authority": "Telangana State Technology Services Limited",
    "turnover_threshold_inr": 50_000_000,   # Rs. 5 Cr
    "turnover_clause": "4.2",
    "experience_years": 5,
    "experience_clause": "4.3",
    "emd_inr": 200_000,                      # Rs. 2,00,000
    "submission_deadline": "2026-03-15",
    "pre_bid_deadline": "2026-02-20",
    "mandatory_docs": [
        "GST Registration Certificate",
        "PAN Card",
        "Audited Balance Sheets (last 3 financial years)",
        "EMD Demand Draft",
        "ISO 9001:2015 Certificate",
    ],
}

_PAGE_1 = f"""TELANGANA STATE TECHNOLOGY SERVICES LIMITED
Hyderabad, Telangana

NOTICE INVITING TENDER

Tender Reference No: {NOTIFICATION_TRUTH['tender_id']}

1. Introduction

1.1 Telangana State Technology Services Limited (hereinafter "the Authority")
invites sealed bids from eligible and qualified bidders for the Supply,
Installation and Commissioning of Network Equipment across twelve district
offices.

1.2 The estimated contract value is Rs. 12,50,00,000 (Rupees Twelve Crore Fifty
Lakh only). Bids shall remain valid for a period of 180 days from the date of
opening.

2. Important Dates

2.1 Last date for submission of pre-bid queries: 20 February 2026, 17:00 hrs IST.
2.2 Last date and time for submission of bids: 15 March 2026, 15:00 hrs IST.
2.3 Technical bids shall be opened on 16 March 2026 at 11:00 hrs IST.

3. Earnest Money Deposit

3.1 Each bidder shall furnish an Earnest Money Deposit of Rs. 2,00,000 (Rupees
Two Lakh only) by way of a Demand Draft drawn in favour of the Authority.
Bids received without the EMD shall be summarily rejected.
"""

_PAGE_2 = """4. Eligibility Criteria

4.1 The bidder shall be a company registered under the Companies Act, 2013, or a
partnership firm registered under the Indian Partnership Act, 1932.

4.2 The bidder shall have an average annual turnover of not less than
Rs. 5 Cr (Rupees Five Crore only) during the last three financial years,
namely 2022-23, 2023-24 and 2024-25. Audited balance sheets shall be furnished
as documentary proof.

4.3 The bidder shall have a minimum of 5 (five) years of experience in the
supply and commissioning of enterprise networking equipment as on the date of
this notification.

4.4 The bidder shall have successfully completed at least two projects of a
value not less than Rs. 2 Cr each within the preceding five years.

4.5 The bidder shall not be blacklisted or debarred by any Central or State
Government department, public sector undertaking, or autonomous body as on the
date of bid submission. A self-declaration to this effect shall be enclosed.
"""

_PAGE_3 = """5. Mandatory Documents

5.1 The following documents shall be enclosed with the technical bid. Bids not
accompanied by all of the documents listed below shall be treated as
non-responsive and shall not be evaluated further.

(i)   GST Registration Certificate
(ii)  PAN Card
(iii) Audited Balance Sheets (last 3 financial years)
(iv)  EMD Demand Draft
(v)   ISO 9001:2015 Certificate

6. Technical Requirements

6.1 All switches supplied shall support IEEE 802.3az Energy Efficient Ethernet.
6.2 The bidder shall provide on-site warranty and support for 3 (three) years.
6.3 The proposed solution shall support centralised management through a single
management console.

7. Submission Format

7.1 Bids shall be submitted in two separate sealed envelopes marked "Technical
Bid" and "Financial Bid" respectively.
7.2 Every page of the bid shall be serially numbered and signed by the
authorised signatory, and shall bear the official seal of the bidder.
7.3 The technical proposal shall not exceed 40 (forty) pages excluding annexures.
"""

_ELIGIBILITY_TABLE = [
    ["Sl. No.", "Criterion", "Requirement", "Clause"],
    ["1", "Average annual turnover", "Not less than Rs. 5 Cr", "4.2"],
    ["2", "Years of experience", "Minimum 5 years", "4.3"],
    ["3", "Similar projects completed", "At least 2, Rs. 2 Cr each", "4.4"],
    ["4", "Blacklisting status", "Not blacklisted / debarred", "4.5"],
]

PAGES = [_PAGE_1, _PAGE_2, _PAGE_3]


def make_notification_pdf(path: str | Path) -> Path:
    """Three-page native-text PDF with a real eligibility table on page 2."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    path = Path(path)
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm
    )
    styles = getSampleStyleSheet()
    story = []

    for index, page_text in enumerate(PAGES):
        for block in page_text.strip().split("\n\n"):
            story.append(Paragraph(block.replace("\n", " "), styles["BodyText"]))
            story.append(Spacer(1, 6))
        if index == 1:  # eligibility summary table on page 2
            story.append(Spacer(1, 10))
            table = Table(_ELIGIBILITY_TABLE)
            # Ruled borders, as real tender eligibility grids have -- this is
            # what pdfplumber's line-based detection keys off.
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ]
                )
            )
            story.append(table)
        if index < len(PAGES) - 1:
            story.append(PageBreak())

    doc.build(story)
    return path


def make_notification_docx(path: str | Path) -> Path:
    """Same content as DOCX, with explicit page breaks."""
    import docx
    from docx.enum.text import WD_BREAK

    path = Path(path)
    document = docx.Document()
    for index, page_text in enumerate(PAGES):
        for block in page_text.strip().split("\n\n"):
            document.add_paragraph(block.replace("\n", " "))
        if index < len(PAGES) - 1:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    table = document.add_table(rows=0, cols=4)
    for row in _ELIGIBILITY_TABLE:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value

    document.save(str(path))
    return path


def make_borderless_table_pdf(path: str | Path) -> Path:
    """A table with no ruling lines -- common in real tenders, and invisible to
    pdfplumber's default line-based table detection."""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table

    path = Path(path)
    SimpleDocTemplate(str(path), pagesize=A4).build([Table(_ELIGIBILITY_TABLE)])
    return path


def make_blank_pdf(path: str | Path, pages: int = 2) -> Path:
    """A PDF with no extractable text -- stands in for a scanned document, so the
    no-OCR degradation path can be tested without shipping a binary scan."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = Path(path)
    pdf = canvas.Canvas(str(path), pagesize=A4)
    for _ in range(pages):
        pdf.showPage()
    pdf.save()
    return path
