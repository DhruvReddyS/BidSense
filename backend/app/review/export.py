"""Portable committee record for the complete company-side review."""

from __future__ import annotations

import io
from datetime import datetime, timezone

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors


def _record(notification, rows) -> tuple[list[str], list[list[str]]]:
    meta = [
        notification.title,
        f"Tender ID: {notification.tender_id}",
        f"Issuing authority: {notification.issuing_authority or 'Not recorded'}",
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
        "Decision notice: This report records system-assisted review. Final selection remains with the evaluation committee.",
    ]
    body = []
    for row in sorted(rows, key=lambda item: (item.status.value, item.vendor_name.casefold())):
        reason = row.elimination_reason or ("Included in the qualified pool." if row.status.value == "shortlisted" else "Cleared Level 1; not selected into the current qualified pool.")
        source = ""
        if row.elimination_clause_ref:
            source = f"Clause {row.elimination_clause_ref}"
            if row.elimination_source_page:
                source += f", page {row.elimination_source_page}"
        body.append([row.vendor_name, row.vendor_id, row.status.value.title(), reason, source])
    return meta, body


def export_docx(notification, rows) -> bytes:
    meta, body = _record(notification, rows)
    document = Document()
    document.add_heading("BidSense Committee Review Record", 0)
    for line in meta:
        document.add_paragraph(line)
    for status in ("Shortlisted", "Pending", "Eliminated"):
        selected = [row for row in body if row[2] == status]
        document.add_heading(status, level=1)
        if not selected:
            document.add_paragraph("No submissions in this state.")
            continue
        for vendor, vendor_id, _, reason, source in selected:
            document.add_heading(vendor, level=2)
            document.add_paragraph(vendor_id)
            document.add_paragraph(reason)
            if source:
                document.add_paragraph(source)
    out = io.BytesIO(); document.save(out); return out.getvalue()


def export_pdf(notification, rows) -> bytes:
    meta, body = _record(notification, rows)
    out = io.BytesIO()
    styles = getSampleStyleSheet()
    story = [Paragraph("BidSense Committee Review Record", styles["Title"]), Spacer(1, 5 * mm)]
    story.extend(Paragraph(line, styles["BodyText"]) for line in meta)
    story.append(Spacer(1, 6 * mm))
    data = [["Bidder", "Status", "Decision record", "Source"]] + [[row[0], row[2], row[3], row[4]] for row in body]
    table = Table(data, colWidths=[38 * mm, 22 * mm, 91 * mm, 27 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24322d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#b8b8b2")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f3ef")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    SimpleDocTemplate(out, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm, bottomMargin=16 * mm).build(story)
    return out.getvalue()
