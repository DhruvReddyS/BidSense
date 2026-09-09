"""Gap report as a document a vendor can hand to their own team (Section 4.6).

A compliance report read on a screen and then quoted from memory is how a
requirement gets missed. The people who actually attach the documents -- the
accounts clerk who produces the audited balance sheet, the director who signs
the power of attorney -- are usually not the person who ran the check, so the
report has to leave the tool.

Both formats render the SAME content from the same `GapReport`, so a vendor
choosing DOCX over PDF does not get a different answer. The renderers share
`_sections()`, which is where the content decisions live; the format modules
below only decide how it looks.

What deliberately survives the export:

  * every clause reference. A to-do without the clause it comes from cannot be
    checked against the tender by whoever receives it.
  * the four-way verdict, including `needs_review` and `not_checked`. Flattening
    those to pass/fail on the way out would undo the whole point of having them.
  * the completion count with its "this is not a score" caveat attached. The
    number travels; so does the sentence that stops it being read as a grade.
  * the data-quality banner, if any. A report built on a thin extraction says so
    on the page, not only in the app.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime, timezone

from app.compliance.models import ActionGroup, CheckStatus, GapReport, Severity

__all__ = ["export_pdf", "export_docx", "ExportContext"]


@dataclass
class ExportContext:
    """Everything on the report that does not live on the GapReport itself."""

    tender_title: str | None = None
    issuing_authority: str | None = None
    submission_deadline: date | None = None
    data_quality_banner: str | None = None
    staleness_banner: str | None = None
    generated_at: datetime = field(default_factory=lambda: datetime.now(timezone.utc))
    extracted_by: str | None = None


from app.compliance.pdf_fonts import pdf_text as _pdf_text


_STATUS_LABEL = {
    CheckStatus.MATCH: "Met",
    CheckStatus.PARTIAL: "Confirm",
    CheckStatus.MISSING: "Missing",
    CheckStatus.MANUAL_CHECK: "Check yourself",
    CheckStatus.NOT_ASSESSABLE: "Could not read",
}

_VERDICT_LABEL = {
    "compliant": "No blocking issues found",
    "needs_review": "No blocking issues, but items need your review",
    "not_compliant": "Blocking issues found — this bid would be rejected as it stands",
    "not_checked": "Nothing could be checked — see data quality below",
}

_GROUP_HEADING = {
    ActionGroup.HARD_FAIL: "Cannot be fixed by attaching a document",
    ActionGroup.UPLOAD: "Documents to attach",
    ActionGroup.CLARIFY: "Values to state clearly in your bid",
    ActionGroup.VERIFY: "To check yourself",
}


def _sections(report: GapReport, context: ExportContext) -> dict:
    """The content, decided once and rendered twice."""
    completion = report.completion
    by_group: dict[ActionGroup, list] = {}
    for action in report.action_list:
        by_group.setdefault(action.group, []).append(action)

    return {
        "title": "Tender compliance check",
        "subtitle": context.tender_title or report.tender_id,
        "meta": [
            ("Tender reference", report.tender_id),
            ("Issuing authority", context.issuing_authority or "not stated"),
            ("Bidder", report.vendor_name or report.vendor_id),
            (
                "Submission deadline",
                context.submission_deadline.strftime("%d %B %Y")
                if context.submission_deadline
                else "not stated",
            ),
            ("Report generated", context.generated_at.strftime("%d %B %Y, %H:%M UTC")),
        ],
        "verdict": _VERDICT_LABEL.get(report.verdict, report.verdict),
        "verdict_key": report.verdict,
        "completion": completion,
        "completion_label": completion.label + (f" - {completion.undetermined} not established either way" if completion.undetermined else ""),
        "banners": [b for b in (context.staleness_banner, context.data_quality_banner) if b],
        "groups": [
            (_GROUP_HEADING[group], by_group[group])
            for group in (
                ActionGroup.HARD_FAIL, ActionGroup.UPLOAD,
                ActionGroup.CLARIFY, ActionGroup.VERIFY,
            )
            if by_group.get(group)
        ],
        "items": report.items,
        "counts": report.counts,
    }


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def export_pdf(report: GapReport, context: ExportContext | None = None) -> bytes:
    from app.compliance.pdf_document import render_pdf
    context = context or ExportContext()
    return render_pdf(report, context, _sections(report, context), _STATUS_LABEL)


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def export_docx(report: GapReport, context: ExportContext | None = None) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    context = context or ExportContext()
    data = _sections(report, context)

    document = Document()
    document.core_properties.title = f"Compliance check — {report.tender_id}"
    document.core_properties.author = "BidSense"

    document.add_heading(data["title"], level=0)
    document.add_paragraph(data["subtitle"])

    table = document.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for key, value in data["meta"]:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = str(value)

    document.add_paragraph()
    verdict = document.add_paragraph()
    run = verdict.add_run(data["verdict"])
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = {
        "not_compliant": RGBColor(0xB4, 0x23, 0x2A),
        "needs_review": RGBColor(0x8A, 0x61, 0x00),
        "compliant": RGBColor(0x1C, 0x6B, 0x45),
    }.get(data["verdict_key"], RGBColor(0x33, 0x33, 0x33))

    completion = data["completion"]
    document.add_paragraph(data["completion_label"]).runs[0].bold = True
    caveat = document.add_paragraph(completion.caveat)
    caveat.runs[0].italic = True
    caveat.runs[0].font.size = Pt(8.5)

    for banner in data["banners"]:
        paragraph = document.add_paragraph(f"Warning: {banner}")
        paragraph.runs[0].font.color.rgb = RGBColor(0x8A, 0x61, 0x00)
        paragraph.runs[0].font.size = Pt(9)

    for heading, actions in data["groups"]:
        document.add_heading(heading, level=2)
        for action in actions:
            bullet = document.add_paragraph(action.action, style="List Bullet")
            if action.clause_ref:
                reference = bullet.add_run(f"  (clause {action.clause_ref})")
                reference.italic = True
                reference.font.size = Pt(8.5)

    document.add_page_break()
    document.add_heading("Every requirement checked", level=1)
    grid = document.add_table(rows=1, cols=5)
    grid.style = "Light Grid Accent 1"
    for index, label in enumerate(("Requirement", "Status", "Required", "Found", "Clause")):
        grid.rows[0].cells[index].text = label
    for item in data["items"]:
        provenance = item.notification_provenance
        cells = grid.add_row().cells
        cells[0].text = item.requirement
        cells[1].text = _STATUS_LABEL.get(item.status, item.status.value)
        cells[2].text = str(item.required_value or "—")
        cells[3].text = str(item.found_value or "—")
        cells[4].text = (
            (f"{provenance.clause_ref or '—'}"
             + (f" p{provenance.source_page}" if provenance.source_page else ""))
            if provenance else "—"
        )

    footer = document.add_paragraph(
        "Generated by BidSense"
        + (f" using {context.extracted_by}" if context.extracted_by else "")
        + ". Format and signing rules are always left to a human reviewer, so no "
          "bid is ever reported as fully clear."
    )
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].italic = True

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
