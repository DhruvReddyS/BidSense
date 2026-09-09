import httpx,json,io,time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from pypdf import PdfReader
from docx import Document
cases=[('DEL/AO-3/RBO-6/AMCC/IN/01','SBI_AMCC_VERSION3'),('TENDER No.01/SE(Electrical)/GHMC/2024-25','VENDOR_supply_01_01')]

def run(case):
 (tender,vendor),fmt=case
 start=time.monotonic()
 r=httpx.post('http://127.0.0.1:8100/api/gap-report/export',params={'fmt':fmt},json={'tender_id':tender,'vendor_id':vendor},timeout=180)
 assert r.status_code==200,(vendor,fmt,r.status_code,r.text[:200])
 if fmt=='pdf':
  doc=PdfReader(io.BytesIO(r.content));text=' '.join(p.extract_text() or '' for p in doc.pages);pages=len(doc.pages)
 else:
  doc=Document(io.BytesIO(r.content));text=' '.join(p.text for p in doc.paragraphs)+' '.join(c.text for t in doc.tables for row in t.rows for c in row.cells);pages=None
 assert 'Warning:' in text and 'Every requirement checked' in text
 return {'vendor':vendor,'format':fmt,'http_status':r.status_code,'bytes':len(r.content),'pages':pages,'seconds':round(time.monotonic()-start,3),'quality_warning_present':True}
with ThreadPoolExecutor(max_workers=4) as pool:results=list(pool.map(run,[(c,f) for c in cases for f in ('pdf','docx')]))
Path('output/pdf-hardening-2026-09-06/live-export.json').write_text(json.dumps(results,indent=2))
print(json.dumps(results,indent=2))
